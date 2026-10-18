"""One-shot Groww XLSX → Pulse OS position book importer.

Usage:
  python -m backend.portfolio.import_groww \\
    --mf "/path/Holdings_Statement_….xlsx" \\
    --stocks "/path/Stocks_Holdings_….xlsx" \\
    --orders "/path/Stocks_Order_History_….xlsx" \\
    --replace
"""

from __future__ import annotations

import argparse
import logging
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from backend.data.nav_client import fetch_latest_nav, search_schemes
from backend.portfolio.ledger import InvestmentLedger

logger = logging.getLogger(__name__)

SECTOR_MAP = {
    # Specific first — generics last
    "pharma": "Pharma",
    "healthcare": "Pharma",
    "automotive": "Auto",
    "auto": "Auto",
    "energy": "Energy",
    "commodities": "Commodities",
    "commodity": "Commodities",
    "consumption": "Consumption",
    "natural resources": "Energy",
    "realty": "Realty",
    "real estate": "Realty",
    "thematic": "Thematic",
    "sectoral": "Sectoral",
}


def _norm(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip())


def _num(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_date(raw: Any) -> Optional[str]:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    s = str(raw).strip()
    for fmt in ("%d-%m-%Y %I:%M %p", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    # e.g. "07-07-2026 09:15 AM"
    m = re.match(r"(\d{2}-\d{2}-\d{4})", s)
    if m:
        try:
            return datetime.strptime(m.group(1), "%d-%m-%Y").date().isoformat()
        except ValueError:
            pass
    return None


def _sector_from_labels(*labels: str) -> str:
    blob = " ".join(_norm(x).lower() for x in labels if x)
    for key, sector in SECTOR_MAP.items():
        if key in blob:
            return sector
    # fallback: first non-empty label title-cased
    for lab in labels:
        if lab:
            return _norm(lab).title()
    return "Other"


def _asset_class_stock(name: str, isin: str = "") -> str:
    u = f"{name} {isin}".upper()
    if "ETF" in u or "BEES" in u or "AMC -" in name.upper() or isin.upper().startswith("INF"):
        # Many Motilal AMC products are ETFs (INF ISINs)
        if "ETF" in u or "BEES" in u or "MOREALTY" in u or "AMC" in name.upper():
            return "etf"
    return "stock"


def resolve_amfi_code(scheme_name: str) -> Optional[str]:
    """Pick best Direct Growth match from mfapi search."""
    # Known gaps in mfapi search coverage (ISIN stored separately by caller if needed)
    known = {
        "whiteoak capital pharma and healthcare fund direct growth": None,
    }
    if scheme_name.lower().strip() in known:
        return known[scheme_name.lower().strip()]

    queries = [scheme_name]
    # Drop trailing boilerplate
    trimmed = re.sub(
        r"\b(direct|plan|growth|option|fund)\b",
        " ",
        scheme_name,
        flags=re.I,
    )
    trimmed = re.sub(r"\s+", " ", trimmed).strip()
    if trimmed and trimmed.lower() != scheme_name.lower():
        queries.append(trimmed)
    words = [w for w in scheme_name.split() if w.lower() not in {"fund", "plan", "the", "and", "of"}]
    if len(words) >= 3:
        queries.append(" ".join(words[:4]))
        queries.append(" ".join(words[:3]))

    results: List[Dict[str, Any]] = []
    seen_codes = set()
    for q in queries:
        for r in search_schemes(q, limit=25):
            code = r["scheme_code"]
            if code in seen_codes:
                continue
            seen_codes.add(code)
            results.append(r)
        if results:
            break

    if not results:
        return None

    name_l = scheme_name.lower()
    scored: List[Tuple[int, str]] = []
    for r in results:
        sn = (r.get("scheme_name") or "").lower()
        score = 0
        if "direct" in sn:
            score += 10
        if "growth" in sn:
            score += 5
        if "idcw" in sn or "dividend" in sn:
            score -= 20
        if "regular" in sn:
            score -= 8
        tokens = [t for t in re.split(r"[^a-z0-9]+", name_l) if len(t) > 2]
        score += sum(2 for t in tokens if t in sn)
        # Require at least some distinctive overlap beyond AMC noise
        distinctive = [t for t in tokens if t not in {"fund", "capital", "india", "direct", "growth", "plan"}]
        hit = sum(1 for t in distinctive if t in sn)
        if distinctive and hit == 0:
            score -= 30
        if sn == name_l or name_l in sn or sn in name_l:
            score += 15
        scored.append((score, r["scheme_code"]))
    scored.sort(key=lambda x: -x[0])
    if not scored or scored[0][0] < 5:
        return None
    return scored[0][1]


def parse_mf_holdings(path: Path) -> Tuple[List[Dict[str, Any]], Optional[str]]:
    """Parse Groww Holdings_Statement XLSX. Returns (rows, as_of_date)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_raw = list(ws.iter_rows(values_only=True))
    wb.close()

    as_of: Optional[str] = None
    header_idx = None
    for i, row in enumerate(rows_raw):
        cells = [_norm(c) for c in row]
        joined = " ".join(cells).lower()
        m = re.search(r"holdings as on\s+(\d{4}-\d{2}-\d{2})", joined)
        if m:
            as_of = m.group(1)
        if cells and cells[0].lower() == "scheme name":
            header_idx = i
            break

    if header_idx is None:
        raise ValueError(f"Could not find Scheme Name header in {path}")

    headers = [_norm(c).lower() for c in rows_raw[header_idx]]
    col = {h: i for i, h in enumerate(headers) if h}

    def get(row, *names, default=None):
        for n in names:
            if n in col and col[n] < len(row):
                return row[col[n]]
        return default

    out: List[Dict[str, Any]] = []
    for row in rows_raw[header_idx + 1 :]:
        if not row or not _norm(row[0]):
            continue
        name = _norm(get(row, "scheme name"))
        if not name or name.lower().startswith("scheme"):
            continue
        units = _num(get(row, "units"))
        invested = _num(get(row, "invested value"))
        current = _num(get(row, "current value"))
        if units <= 0 and invested <= 0:
            continue
        category = _norm(get(row, "category"))
        sub = _norm(get(row, "sub-category", "sub category"))
        nav = (current / units) if units else 0
        avg = (invested / units) if units else 0
        sector = _sector_from_labels(sub, category, name)
        out.append(
            {
                "instrument": name,
                "asset_class": "mf",
                "units": units,
                "invested_amount": invested,
                "avg_cost": avg,
                "last_nav": nav,
                "market_value": current,
                "sector": sector,
                "category": category,
                "sub_category": sub,
                "folio": _norm(get(row, "folio no.", "folio no")),
            }
        )
    return out, as_of


def parse_stock_holdings(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_raw = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = None
    for i, row in enumerate(rows_raw):
        cells = [_norm(c).lower() for c in row]
        if cells and cells[0] == "stock name":
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(f"Could not find Stock Name header in {path}")

    headers = [_norm(c).lower() for c in rows_raw[header_idx]]
    col = {h: i for i, h in enumerate(headers) if h}

    def get(row, *names):
        for n in names:
            if n in col and col[n] < len(row):
                return row[col[n]]
        return None

    out: List[Dict[str, Any]] = []
    for row in rows_raw[header_idx + 1 :]:
        if not row or not _norm(row[0]):
            continue
        name = _norm(get(row, "stock name"))
        if not name:
            continue
        qty = _num(get(row, "quantity"))
        avg = _num(get(row, "average buy price", "average buy price"))
        buy_value = _num(get(row, "buy value"))
        close = _num(get(row, "closing price"))
        close_value = _num(get(row, "closing value"))
        isin = _norm(get(row, "isin"))
        if qty <= 0:
            continue
        out.append(
            {
                "instrument": name,
                "asset_class": _asset_class_stock(name, isin),
                "units": qty,
                "invested_amount": buy_value or (qty * avg),
                "avg_cost": avg,
                "last_nav": close or avg,
                "market_value": close_value or (qty * (close or avg)),
                "scheme_code": isin,
                "sector": _sector_from_labels(name),
                "symbol": name,
            }
        )
    return out


def parse_stock_orders(path: Path) -> Dict[str, str]:
    """Return {isin_or_symbol_upper: earliest_entry_date ISO}."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_raw = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = None
    for i, row in enumerate(rows_raw):
        cells = [_norm(c).lower() for c in row]
        if cells and cells[0] in ("stock name", "stockname"):
            header_idx = i
            break
    if header_idx is None:
        return {}

    headers = [_norm(c).lower() for c in rows_raw[header_idx]]
    col = {h: i for i, h in enumerate(headers) if h}

    def get(row, *names):
        for n in names:
            if n in col and col[n] < len(row):
                return row[col[n]]
        return None

    earliest: Dict[str, str] = {}
    for row in rows_raw[header_idx + 1 :]:
        if not row:
            continue
        status = _norm(get(row, "order status")).lower()
        tx_type = _norm(get(row, "type")).upper()
        if status and status not in ("executed", "complete", "completed"):
            continue
        if tx_type and tx_type not in ("BUY", "B"):
            continue
        d = _parse_date(get(row, "execution date and time", "execution date"))
        if not d:
            continue
        isin = _norm(get(row, "isin")).upper()
        symbol = _norm(get(row, "symbol")).upper()
        name = _norm(get(row, "stock name")).upper()
        for key in (isin, symbol, name):
            if not key:
                continue
            if key not in earliest or d < earliest[key]:
                earliest[key] = d
    return earliest


def _position_payload(row: Dict[str, Any], entry_date: str, scheme_code: Optional[str] = None) -> Dict[str, Any]:
    sector = row.get("sector") or "Other"
    today = date.today()
    return {
        "instrument": row["instrument"],
        "asset_class": row["asset_class"],
        "scheme_code": scheme_code if scheme_code is not None else row.get("scheme_code"),
        "sector": sector,
        "sleeve": "satellite",
        "units": row["units"],
        "invested_amount": row["invested_amount"],
        "avg_cost": row.get("avg_cost") or 0,
        "entry_date": entry_date,
        "last_nav": row.get("last_nav"),
        "last_nav_date": entry_date,
        "thesis": (
            f"Satellite tilt: {sector} — seeded from Groww holdings. "
            "Refine thesis and exit rules in review."
        ),
        "horizon_months_lo": 6,
        "horizon_months_hi": 9,
        "review_date": (today + timedelta(days=30)).isoformat(),
        "exit_by_date": (today + timedelta(days=9 * 30)).isoformat(),
        "stop_loss_pct": 12,
        "notes": f"Imported from Groww XLSX on {today.isoformat()}",
    }


def import_groww(
    mf_path: Optional[Path] = None,
    stocks_path: Optional[Path] = None,
    orders_path: Optional[Path] = None,
    replace: bool = False,
    resolve_amfi: bool = True,
    refresh_nav: bool = True,
) -> Dict[str, Any]:
    ledger = InvestmentLedger()
    cleared = None
    if replace:
        cleared = ledger.clear_book()
        logger.info("Cleared book: %s", cleared)

    created: List[Dict[str, Any]] = []
    amfi_resolved = 0
    nav_refreshed = 0

    if mf_path:
        mf_rows, as_of = parse_mf_holdings(mf_path)
        entry = as_of or date.today().isoformat()
        for row in mf_rows:
            code = None
            if resolve_amfi:
                code = resolve_amfi_code(row["instrument"])
                if code:
                    amfi_resolved += 1
                    logger.info("AMFI %s → %s", row["instrument"][:50], code)
                else:
                    # mfapi may not index newer schemes; keep ISIN when known
                    isin_fallback = {
                        "whiteoak capital pharma and healthcare fund direct growth": "INF03VN01886",
                    }
                    code = isin_fallback.get(row["instrument"].lower().strip())
                    if code:
                        logger.info("ISIN fallback %s → %s", row["instrument"][:40], code)
                    else:
                        logger.warning("No AMFI code for %s", row["instrument"])
            pos = ledger.create_position(_position_payload(row, entry, code))
            if refresh_nav and code and code.isdigit():
                info = fetch_latest_nav(code)
                if info:
                    pos = ledger.update_nav(pos["id"], info["nav"], info.get("date")) or pos
                    nav_refreshed += 1
            created.append(pos)

    order_dates: Dict[str, str] = {}
    if orders_path and orders_path.exists():
        order_dates = parse_stock_orders(orders_path)

    if stocks_path:
        stock_rows = parse_stock_holdings(stocks_path)
        for row in stock_rows:
            isin = (row.get("scheme_code") or "").upper()
            name_u = row["instrument"].upper()
            entry = (
                order_dates.get(isin)
                or order_dates.get(name_u)
                or next((v for k, v in order_dates.items() if k and k in name_u), None)
                or date.today().isoformat()
            )
            pos = ledger.create_position(_position_payload(row, entry, row.get("scheme_code")))
            created.append(pos)

    summary = ledger.portfolio_summary()
    return {
        "cleared": cleared,
        "created_count": len(created),
        "amfi_resolved": amfi_resolved,
        "nav_refreshed": nav_refreshed,
        "positions": [
            {
                "id": p["id"],
                "instrument": p["instrument"],
                "asset_class": p["asset_class"],
                "sector": p["sector"],
                "units": p["units"],
                "invested_amount": p["invested_amount"],
                "market_value": p["market_value"],
                "scheme_code": p.get("scheme_code"),
            }
            for p in created
        ],
        "summary": summary,
    }


def main(argv: Optional[List[str]] = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    p = argparse.ArgumentParser(description="Seed Pulse OS from Groww XLSX exports")
    p.add_argument("--mf", type=Path, help="Holdings_Statement_*.xlsx (mutual funds)")
    p.add_argument("--stocks", type=Path, help="Stocks_Holdings_Statement_*.xlsx")
    p.add_argument("--orders", type=Path, help="Stocks_Order_History_*.xlsx")
    p.add_argument("--replace", action="store_true", help="Clear existing positions first")
    p.add_argument("--no-amfi", action="store_true", help="Skip AMFI scheme code lookup")
    p.add_argument("--no-nav-refresh", action="store_true", help="Skip live NAV refresh")
    args = p.parse_args(argv)

    if not args.mf and not args.stocks:
        p.error("Provide at least --mf or --stocks")

    for path, label in ((args.mf, "mf"), (args.stocks, "stocks"), (args.orders, "orders")):
        if path and not path.exists():
            p.error(f"{label} file not found: {path}")

    result = import_groww(
        mf_path=args.mf,
        stocks_path=args.stocks,
        orders_path=args.orders,
        replace=args.replace,
        resolve_amfi=not args.no_amfi,
        refresh_nav=not args.no_nav_refresh,
    )
    print(f"Created {result['created_count']} positions")
    if result["cleared"]:
        print(f"Cleared: {result['cleared']}")
    print(f"AMFI resolved: {result['amfi_resolved']} · NAV refreshed: {result['nav_refreshed']}")
    s = result["summary"]
    print(
        f"Book: invested ₹{s['total_invested']:,.2f} · value ₹{s['total_value']:,.2f} · "
        f"P&L {s['pnl_pct']:+.2f}% · {s['position_count']} open"
    )
    for pos in result["positions"]:
        print(
            f"  - [{pos['asset_class']}] {pos['instrument'][:55]} | "
            f"{pos['units']} u | ₹{pos['market_value']:,.2f} | code={pos.get('scheme_code') or '—'}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
